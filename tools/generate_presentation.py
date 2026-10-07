#!/usr/bin/env python3
"""Generate the editable AIORI-3 hackathon presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "assets" / "AIORI-3-Hackathon-Presentation.pptx"

NAVY = RGBColor(8, 17, 31)
PANEL = RGBColor(17, 31, 50)
PANEL_LIGHT = RGBColor(24, 43, 65)
WHITE = RGBColor(239, 246, 255)
MUTED = RGBColor(156, 176, 200)
CYAN = RGBColor(83, 216, 232)
GREEN = RGBColor(101, 230, 177)
AMBER = RGBColor(255, 198, 109)
RED = RGBColor(255, 133, 133)

SW, SH = 13.333, 7.5


def rgb_fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def textbox(slide, x, y, w, h, text, size=18, color=WHITE, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, color=PANEL, radius=True, outline=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    rgb_fill(shape, color)
    if outline:
        shape.line.color.rgb = outline
        shape.line.width = Pt(1)
    return shape


def base_slide(prs, index, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    textbox(slide, 0.65, 0.28, 9.8, 0.28, "AIORI-3  /  LAST MINUTE CODERS",
            11, CYAN, True)
    textbox(slide, 11.5, 0.28, 1.15, 0.28, f"{index:02d} / 08",
            10, MUTED, True, align=PP_ALIGN.RIGHT)
    textbox(slide, 0.65, 0.78, 12.0, 0.65, title, 28, WHITE, True)
    if subtitle:
        textbox(slide, 0.67, 1.42, 11.9, 0.45, subtitle, 14, MUTED)
    rect(slide, 0.65, 7.08, 12.0, 0.02, PANEL_LIGHT, radius=False)
    textbox(slide, 0.65, 7.14, 9.0, 0.2,
            "PRIVACY-AWARE IPv6 IoT SECURITY MONITORING", 8, MUTED, True)
    textbox(slide, 11.75, 7.14, 0.9, 0.2, "AIORI-3", 8, CYAN, True,
            align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, heading, detail, accent=CYAN):
    rect(slide, x, y, w, h, PANEL, outline=PANEL_LIGHT)
    rect(slide, x + 0.18, y + 0.22, 0.06, h - 0.44, accent, radius=False)
    textbox(slide, x + 0.4, y + 0.18, w - 0.58, 0.43,
            heading, 17, WHITE, True)
    textbox(slide, x + 0.4, y + 0.72, w - 0.58, h - 0.87,
            detail, 12, MUTED, valign=MSO_ANCHOR.TOP)


def bullet_list(slide, x, y, w, items, line_height=0.68, size=16):
    for idx, item in enumerate(items):
        yy = y + idx * line_height
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(yy + 0.15), Inches(0.11), Inches(0.11)
        )
        rgb_fill(dot, CYAN if idx % 2 == 0 else GREEN)
        textbox(slide, x + 0.26, yy, w - 0.26, line_height - 0.02,
                item, size, WHITE, valign=MSO_ANCHOR.TOP)


def connector(slide, x1, y1, x2, y2):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = CYAN
    line.line.width = Pt(2)
    line.line.end_arrowhead = True


def make_deck():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    prs.core_properties.title = (
        "Privacy-Aware IPv6 IoT Security Monitoring Across Address Rotation"
    )
    prs.core_properties.subject = "AIORI-3 Hackathon presentation"
    prs.core_properties.author = "Last Minute Coders"

    # 1. Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    rect(slide, 8.65, 0.0, 4.68, SH, PANEL, radius=False)
    textbox(slide, 0.8, 0.65, 7.5, 0.45, "AIORI-3 HACKATHON  |  PS007",
            13, CYAN, True)
    textbox(slide, 0.8, 1.55, 7.55, 2.05,
            "Privacy-Aware IPv6 IoT Security Monitoring",
            31, WHITE, True, valign=MSO_ANCHOR.TOP)
    textbox(slide, 0.82, 3.68, 7.3, 0.72,
            "Across Address Rotation", 24, GREEN, True)
    textbox(slide, 0.82, 4.82, 7.3, 0.48,
            "Last Minute Coders", 19, WHITE, True)
    textbox(slide, 0.82, 5.35, 7.3, 0.4,
            "Akrit Goyal  ·  Vedika Pathak", 14, MUTED)
    textbox(slide, 0.82, 6.02, 7.3, 0.35,
            "Faculty Mentor: Sachin Kumar", 12, MUTED)
    textbox(slide, 9.15, 1.25, 3.35, 0.42, "PROBLEM STATEMENT", 11, CYAN, True)
    textbox(slide, 9.15, 1.78, 3.3, 0.68, "A3-PS007-TC236", 23, WHITE, True)
    card(slide, 9.12, 2.8, 3.45, 1.15, "Detect", "Repeated IPv6 address changes", CYAN)
    card(slide, 9.12, 4.15, 3.45, 1.15, "Protect", "Pseudonymous operator views", GREEN)
    textbox(slide, 9.15, 6.35, 3.35, 0.35, "AIORI-3  |  OCTOBER 2026", 10, MUTED, True)

    # 2. Problem
    slide = base_slide(prs, 2, "The monitoring gap",
                       "IPv6 privacy rotation is useful, but it can make suspicious behavior harder to correlate.")
    card(slide, 0.75, 2.15, 3.75, 1.7, "Addresses change",
         "IoT devices may use changing IPv6 addresses as networks and privacy settings evolve.", AMBER)
    card(slide, 4.8, 2.15, 3.75, 1.7, "Correlation gets harder",
         "Monitoring that assumes a stable address can miss a repeated change pattern.", CYAN)
    card(slide, 8.85, 2.15, 3.75, 1.7, "Raw data adds risk",
         "Full device identifiers and addresses can expose sensitive operational details.", RED)
    rect(slide, 0.75, 4.45, 11.85, 1.35, PANEL_LIGHT)
    textbox(slide, 1.05, 4.62, 11.2, 0.37, "Our design goal", 16, GREEN, True)
    textbox(slide, 1.05, 5.03, 11.15, 0.55,
            "Detect repeated IPv6 changes while limiting identity and address details in operator-facing responses.",
            19, WHITE, True)

    # 3. Solution
    slide = base_slide(prs, 3, "The prototype: correlate privately",
                       "A lightweight event-to-alert workflow for a clear, repeatable hackathon demonstration.")
    card(slide, 0.8, 2.15, 3.65, 1.55, "1. Ingest telemetry",
         "FastAPI accepts a device ID, IPv6 address, event type, and network context.", CYAN)
    card(slide, 4.85, 2.15, 3.65, 1.55, "2. Pseudonymize",
         "HMAC-based keys correlate the device and address fingerprint.", GREEN)
    card(slide, 8.9, 2.15, 3.65, 1.55, "3. Flag rotation",
         "Two or more consecutive address changes in 15 minutes create an alert.", AMBER)
    connector(slide, 4.5, 2.92, 4.8, 2.92)
    connector(slide, 8.55, 2.92, 8.85, 2.92)
    rect(slide, 0.8, 4.25, 11.75, 1.55, PANEL)
    textbox(slide, 1.08, 4.45, 2.0, 0.35, "OPERATOR VIEW", 12, CYAN, True)
    textbox(slide, 1.08, 4.95, 11.0, 0.52,
            "A partial pseudonym + severity + explainable alert. No raw device ID or full IPv6 address in summary/alert responses.",
            17, WHITE, True)

    # 4. Architecture
    slide = base_slide(prs, 4, "How the pieces fit",
                       "A small Python service keeps the demo deployable without external infrastructure.")
    nodes = [
        (0.8, "IoT / Demo\nEvents", "Sample telemetry", CYAN),
        (3.25, "FastAPI\nIngestion", "POST /api/v1/events", GREEN),
        (5.7, "Privacy +\nHeuristics", "HMAC + rotation window", AMBER),
        (8.15, "SQLite\nPersistence", "Events + alert records", CYAN),
        (10.6, "Summary /\nAlerts API", "GET /summary, /alerts", GREEN),
    ]
    for x, heading, detail, accent in nodes:
        rect(slide, x, 2.85, 1.9, 1.55, PANEL, outline=PANEL_LIGHT)
        textbox(slide, x + 0.12, 3.05, 1.66, 0.65, heading, 15, accent, True, align=PP_ALIGN.CENTER)
        textbox(slide, x + 0.1, 3.82, 1.7, 0.35, detail, 10, MUTED, align=PP_ALIGN.CENTER)
    for x in (2.77, 5.22, 7.67, 10.12):
        connector(slide, x, 3.62, x + 0.42, 3.62)
    textbox(slide, 1.0, 5.05, 11.0, 0.45,
            "Python  |  FastAPI  |  SQLAlchemy  |  SQLite  |  Docker-ready",
            16, WHITE, True, align=PP_ALIGN.CENTER)

    # 5. Live demonstration
    slide = base_slide(prs, 5, "Live demo: from events to alert",
                       "Use documentation-only IPv6 addresses (2001:db8::/32) and a fresh demo database.")
    demo_steps = [
        ("01", "Start service", "Check /health"),
        ("02", "Send event 1", "2001:db8:1::1"),
        ("03", "Rotate twice", "::2, then ::3"),
        ("04", "Review result", "/api/v1/summary"),
    ]
    for idx, (number, heading, detail) in enumerate(demo_steps):
        x = 0.85 + idx * 3.1
        rect(slide, x, 2.5, 2.65, 1.8, PANEL, outline=PANEL_LIGHT)
        textbox(slide, x + 0.22, 2.72, 0.7, 0.35, number, 13, CYAN, True)
        textbox(slide, x + 0.22, 3.15, 2.2, 0.4, heading, 16, WHITE, True)
        textbox(slide, x + 0.22, 3.65, 2.2, 0.4, detail, 12, MUTED)
        if idx < 3:
            connector(slide, x + 2.7, 3.4, x + 3.02, 3.4)
    rect(slide, 0.85, 4.85, 11.95, 1.05, PANEL_LIGHT)
    textbox(slide, 1.12, 5.02, 11.3, 0.62,
            "Expected behavior: after two distinct address changes in the 15-minute window, the API returns a medium-severity rotation alert.",
            15, GREEN, True)

    # 6. Privacy boundaries
    slide = base_slide(prs, 6, "Privacy controls and honest boundaries",
                       "Privacy-aware does not mean the prototype never stores sensitive data.")
    card(slide, 0.85, 2.0, 5.55, 1.45, "What the API limits",
         "Device IDs become HMAC pseudonyms. Summary and alert responses omit full IPv6 addresses.", GREEN)
    card(slide, 6.85, 2.0, 5.55, 1.45, "What is stored internally",
         "The demo database retains raw IPv6 values alongside pseudonymous keys for event analysis.", AMBER)
    card(slide, 0.85, 3.85, 5.55, 1.45, "Prototype scope",
         "Current detection is a transparent consecutive-address-change heuristic with a 15-minute window.", CYAN)
    card(slide, 6.85, 3.85, 5.55, 1.45, "Not yet implemented",
         "Reconnect-loop and network-segment anomaly detection are future work, not current signals.", RED)
    textbox(slide, 0.9, 5.85, 11.6, 0.45,
            "Production use would require protected secrets, access controls, retention policies, and validation.",
            14, MUTED, True)

    # 7. Impact / next steps
    slide = base_slide(prs, 7, "Why it matters - and what comes next",
                       "Make rotating IPv6 behavior easier to review without turning the prototype into a black box.")
    bullet_list(slide, 0.95, 2.15, 5.2, [
        "Operator-readable reason for the alert",
        "Compact FastAPI + SQLite deployment",
        "Pseudonymous correlation for demo events",
        "Clear threshold that can be tuned and tested",
    ], line_height=0.72, size=15)
    rect(slide, 6.75, 2.05, 5.65, 3.5, PANEL)
    textbox(slide, 7.05, 2.32, 4.9, 0.4, "NEXT ITERATIONS", 13, CYAN, True)
    bullet_list(slide, 7.05, 2.95, 4.95, [
        "Add authenticated operator access",
        "Define safe retention and deletion",
        "Evaluate heuristics on representative data",
        "Explore reconnect and segment signals",
    ], line_height=0.55, size=13)
    textbox(slide, 0.95, 5.75, 11.4, 0.4,
            "A demonstrable starting point for privacy-aware IoT monitoring - not a claim of production readiness.",
            15, GREEN, True)

    # 8. Closing
    slide = base_slide(prs, 8, "Thank you",
                       "Privacy-Aware IPv6 IoT Security Monitoring Across Address Rotation")
    textbox(slide, 0.95, 2.25, 6.4, 0.6, "Last Minute Coders", 27, WHITE, True)
    textbox(slide, 0.98, 3.03, 6.0, 0.45, "Akrit Goyal  ·  Vedika Pathak", 18, MUTED)
    textbox(slide, 0.98, 3.58, 6.0, 0.38, "Faculty Mentor: Sachin Kumar", 14, MUTED)
    rect(slide, 7.8, 2.05, 4.35, 2.5, PANEL, outline=PANEL_LIGHT)
    textbox(slide, 8.1, 2.35, 3.75, 0.38, "PROBLEM STATEMENT", 12, CYAN, True)
    textbox(slide, 8.1, 2.88, 3.75, 0.55, "A3-PS007-TC236", 23, WHITE, True)
    textbox(slide, 8.1, 3.68, 3.75, 0.4, "Questions & discussion", 15, GREEN, True)
    textbox(slide, 0.98, 5.55, 11.2, 0.5,
            "Repository: github.com/goyalakrit-cyber/ps-007-privacy-aware-ipv6-iot-security-monitoring",
            11, MUTED)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f"Created presentation: {OUTPUT}")


if __name__ == "__main__":
    make_deck()
